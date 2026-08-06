"""Generate charts for Excel files.
Configure chart type and range in CONFIG."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.utils.cell import range_boundaries


CONFIG = {
    "input_file": "C:\\Users\\xxx\\Desktop\\test.xlsx",
    "output_file": "C:\\Users\\xxx\\Desktop\\Test_Chart.xlsx",
    "sheet_name": "Sheet1",
    "charts": [
        {
            "type": "line",
            "title": "Sheet1",
            "data_range": "B1:B12",
            "categories_range": "A1:A12",
            "titles_from_data": False,
            "position": "D2",
        }
    ],
}


def build_chart(chart_type: str):
    if chart_type == "bar":
        return BarChart()
    if chart_type == "line":
        return LineChart()
    if chart_type == "pie":
        return PieChart()
    raise ValueError(f"Unsupported chart type: {chart_type}")


def make_reference(sheet, cell_range: str) -> Reference:
    min_col, min_row, max_col, max_row = range_boundaries(cell_range)
    if min_col is None or min_row is None or max_col is None or max_row is None:
        raise ValueError(f"Invalid cell range: {cell_range}")
    return Reference(
        sheet,
        min_col=int(min_col),
        min_row=int(min_row),
        max_col=int(max_col),
        max_row=int(max_row),
    )


def main() -> None:
    workbook = load_workbook(CONFIG["input_file"])
    sheet = workbook[CONFIG["sheet_name"]]

    for chart_cfg in CONFIG["charts"]:
        chart = build_chart(chart_cfg["type"])
        chart.title = chart_cfg.get("title")

        data_ref = make_reference(sheet, chart_cfg["data_range"])
        titles_from_data = chart_cfg.get("titles_from_data", True)
        chart.add_data(data_ref, titles_from_data=titles_from_data)

        if chart_cfg.get("categories_range"):
            categories_ref = make_reference(sheet, chart_cfg["categories_range"])
            chart.set_categories(categories_ref)

        position = chart_cfg.get("position")
        if position:
            chart.anchor = position
        sheet.add_chart(chart)

    output_path = Path(CONFIG["output_file"])
    workbook.save(output_path)
    print(f"Chart generation completed: {output_path}")


if __name__ == "__main__":
    main()
